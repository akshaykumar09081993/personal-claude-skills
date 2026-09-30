#!/usr/bin/env python3
"""build_orders_xlsx.py — turn dl_orders.js output (otto-orders.json) into an Excel workbook.

Usage:
    python3 build_orders_xlsx.py [in.json] [out.xlsx]
    # defaults: ./otto-orders.json  ->  ./OTTO_Orders.xlsx
Requires: openpyxl  (pip install openpyxl)

Consumes the JSON written by scripts/dl_orders.js, i.e. records shaped like:
  { route, capturedAt, customers:[{id,name,addr}], weeks:[..],
    records:[ { customerId, customerName, week, days:[7 "Mmm-D Ddd"],
               products:[{name,sku,tf,rtn4wk,fo:[7],so:[7],weekTotalFO}],
               grandDays:[7], verify:{colSums,ok,mismatches} } ] }

Sheets produced:
  - Summary:      route, capture time, store list, week range, record count.
  - All Orders:   flat long table (Store, Week, Product, SKU, TF, 4wkRtn%, Sun..Sat, WeekTotal) — pivot-ready.
  - One per store: weekly blocks (product rows x day columns) with a computed TOTAL row, OTTO's
                   Grand Total row, and a Match? row for verification.
  - Verification: every store x week, column-sum vs OTTO grand total, OK / MISMATCH.
"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.getcwd(), 'otto-orders.json')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.getcwd(), 'OTTO_Orders.xlsx')

data = json.load(open(SRC))
records = data.get('records', [])

WK = ['Sun', 'Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat']
def weekday(label):  # "Sep-27 Sun" -> "Sun"
    return label.split()[-1] if label else ''

def product_codes(pr):
    """Return (material_code, description, upc, bimbo_code) for a product.
    name = "1290542 DEMPSTERS HOLSUM BRD WW 3X570G"  (leading material code)
    sku  = "068721210146 - 921964"                   (UPC - Bimbo article code)
    Prefers explicit fields (code/upc/bimboCode) if the scraper provided them."""
    name = pr.get('name', '') or ''
    parts = name.split(' ', 1)
    if parts and parts[0].isdigit():
        material, desc = parts[0], (parts[1] if len(parts) > 1 else '')
    else:
        material, desc = '', name
    sku = pr.get('sku', '') or ''
    if ' - ' in sku:
        upc, bimbo = [s.strip() for s in sku.split(' - ', 1)]
    else:
        upc, bimbo = sku.strip(), ''
    return pr.get('code', material), desc, pr.get('upc', upc), pr.get('bimboCode', bimbo)

hdr_font = Font(bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='305496')
tot_fill = PatternFill('solid', fgColor='FCE4D6')
gt_fill  = PatternFill('solid', fgColor='E2EFDA')
bad_fill = PatternFill('solid', fgColor='FFC7CE')
ok_fill  = PatternFill('solid', fgColor='C6EFCE')
bold = Font(bold=True)
center = Alignment(horizontal='center')
thin = Side(style='thin', color='D9D9D9')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = center; cell.border = border

wb = Workbook()

# ---------------- Summary ----------------
ws = wb.active; ws.title = 'Summary'
ws['A1'] = 'OTTO Order Export'; ws['A1'].font = Font(bold=True, size=16)
ws['A3'] = 'Route';        ws['B3'] = data.get('route')
ws['A4'] = 'Captured';     ws['B4'] = data.get('capturedAt')
ws['A5'] = 'Weeks';        ws['B5'] = ', '.join(data.get('weeks', []))
ws['A6'] = 'Stores';       ws['B6'] = len(data.get('customers', []))
ws['A7'] = 'Records';      ws['B7'] = len(records)
for r in range(3, 8): ws.cell(row=r, column=1).font = bold
ws['A9'] = 'Stores'; ws['A9'].font = bold
ws['A10'] = 'ID'; ws['B10'] = 'Name'; ws['C10'] = 'Address'
style_header(ws, 10, 3)
for i, c in enumerate(data.get('customers', [])):
    ws.cell(row=11 + i, column=1, value=c.get('id'))
    ws.cell(row=11 + i, column=2, value=c.get('name'))
    ws.cell(row=11 + i, column=3, value=c.get('addr'))
ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 38; ws.column_dimensions['C'].width = 30

# ---------------- All Orders (flat) ----------------
# "Code" = Bimbo article code (always present, 2nd part of the sku). UPC = barcode. Material code
# (leading number in the product name) exists only for some stores, so it's an extra trailing column.
fs = wb.create_sheet('All Orders')
cols = ['Store ID', 'Store', 'Week', 'Code', 'Product', 'UPC', 'TF', '4wk Rtn%'] + WK + ['Week Total', 'Material Code']
fs.append(cols); style_header(fs, 1, len(cols))
for rec in records:
    wdays = [weekday(d) for d in rec['days']]
    idx = {wdays[i]: i for i in range(len(wdays))}
    for pr in rec['products']:
        fo = pr['fo']
        material, desc, upc, bimbo = product_codes(pr)
        row = [rec['customerId'], rec['customerName'], rec['week'], bimbo, desc, upc, pr.get('tf', ''), pr.get('rtn4wk', '')]
        row += [fo[idx[d]] if d in idx and idx[d] < len(fo) else 0 for d in WK]
        row += [pr.get('weekTotalFO', sum(fo)), material]
        fs.append(row)
fs.freeze_panes = 'A2'
widths = [12, 34, 6, 10, 38, 14, 6, 9] + [6] * 7 + [11, 13]
for i, w in enumerate(widths): fs.column_dimensions[get_column_letter(i + 1)].width = w

# ---------------- Per-store sheets ----------------
def safe_title(name, used):
    t = ''.join(ch for ch in name if ch not in '[]:*?/\\')[:28]
    base = t; n = 1
    while t in used: n += 1; t = f'{base[:25]}~{n}'
    used.add(t); return t

used_titles = set()
by_store = {}
for rec in records:
    by_store.setdefault((rec['customerId'], rec['customerName']), []).append(rec)

for (cid, cname), recs in by_store.items():
    ws = wb.create_sheet(safe_title(f'{cid} {cname}', used_titles))
    ws['A1'] = f'{cid} — {cname}'; ws['A1'].font = Font(bold=True, size=13)
    r = 3
    recs.sort(key=lambda x: int(x['week']))
    D0 = 3  # first day column (col 1 = Code, col 2 = Product)
    for rec in recs:
        ndays = len(rec['days'])
        ws.cell(row=r, column=1, value=f"Week {rec['week']}").font = bold
        r += 1
        head = ['Code', 'Product'] + rec['days'] + ['Week Total']
        for j, h in enumerate(head):
            ws.cell(row=r, column=1 + j, value=h)
        style_header(ws, r, len(head))
        r += 1
        col_sums = [0] * ndays
        for pr in rec['products']:
            material, desc, upc, bimbo = product_codes(pr)
            ws.cell(row=r, column=1, value=bimbo)   # Bimbo article code (always present)
            ws.cell(row=r, column=2, value=desc)
            for j, v in enumerate(pr['fo']):
                ws.cell(row=r, column=D0 + j, value=v)
                if j < len(col_sums): col_sums[j] += v
            ws.cell(row=r, column=D0 + ndays, value=pr.get('weekTotalFO', sum(pr['fo'])))
            r += 1
        ws.cell(row=r, column=1, value='TOTAL (sum of column)').font = bold
        for j, s in enumerate(col_sums):
            c = ws.cell(row=r, column=D0 + j, value=s); c.font = bold; c.fill = tot_fill
        ws.cell(row=r, column=D0 + ndays, value=sum(col_sums)).font = bold
        r += 1
        gd = rec.get('grandDays', [])
        ws.cell(row=r, column=1, value='OTTO Grand Total').font = bold
        for j in range(ndays):
            v = gd[j] if j < len(gd) else None
            c = ws.cell(row=r, column=D0 + j, value=v); c.font = bold; c.fill = gt_fill
        r += 1
        ws.cell(row=r, column=1, value='Match?').font = bold
        for j in range(ndays):
            cs = col_sums[j]; gv = gd[j] if j < len(gd) else None
            ok = (gv is not None and cs == gv)
            c = ws.cell(row=r, column=D0 + j, value='OK' if ok else 'DIFF')
            c.fill = ok_fill if ok else bad_fill; c.alignment = center
        r += 2
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 38
    for j in range(3, 13): ws.column_dimensions[get_column_letter(j)].width = 11

# ---------------- Verification ----------------
vs = wb.create_sheet('Verification')
vcols = ['Store ID', 'Store', 'Week'] + [f'{d} col' for d in WK] + [f'{d} GT' for d in WK] + ['Result']
vs.append(vcols); style_header(vs, 1, len(vcols))
allok = True
for rec in records:
    wdays = [weekday(d) for d in rec['days']]
    idx = {wdays[i]: i for i in range(len(wdays))}
    v = rec.get('verify', {})
    cs = v.get('colSums', [])
    gd = rec.get('grandDays', [])
    ordered_cs = [cs[idx[d]] if d in idx and idx[d] < len(cs) else '' for d in WK]
    ordered_gt = [gd[idx[d]] if d in idx and idx[d] < len(gd) else '' for d in WK]
    ok = v.get('ok', False)
    allok = allok and ok
    row = [rec['customerId'], rec['customerName'], rec['week']] + ordered_cs + ordered_gt + ['OK' if ok else 'MISMATCH']
    vs.append(row)
    rc = vs.cell(row=vs.max_row, column=len(vcols))
    rc.fill = ok_fill if ok else bad_fill; rc.font = bold
vs.freeze_panes = 'A2'
for i in range(len(vcols)): vs.column_dimensions[get_column_letter(i + 1)].width = 10
vs.column_dimensions['B'].width = 34

wb.save(OUT)
print('wrote', OUT)
print('records:', len(records), 'all verified OK:', allok)
